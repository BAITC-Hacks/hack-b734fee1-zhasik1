"""Read-only supplier adapters. No extraction, source edits, or formula execution."""
from dataclasses import dataclass, field
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
from zipfile import ZipFile
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from qor.contracts import KEY, SNAPSHOT

ALIASES = {
    "sku": ["sku", "код", "код 1с", "код 1c", "номенклатура.код"],
    "unit": ["unit", "ед", "ед.", "ед. изм.", "единица"],
    "quantity": ["quantity", "qty", "количество", "кол-во"],
    "date": ["date", "дата"], "document": ["document", "номер"],
    "warehouse": ["warehouse", "склад"], "name": ["name", "наименование", "номенклатура"],
    "article": ["article", "артикул", "артикул поставщика", "артикул иэк"],
    "free_stock": ["free_stock", "свободный остаток", "доступный остаток"],
    "reserved": ["reserved", "резерв"], "as_of": ["as_of", "дата остатка"],
    "eta": ["eta", "дата поставки", "дата прихода"], "inbound_qty": ["inbound_qty", "в пути"],
    "moq": ["moq", "min_order_qty", "минимальный заказ", "мин. разр. к отгр."],
    "pack_multiple": ["pack_multiple", "кратность"], "category": ["category", "категория", "категория товара"],
    "client_id_anonymized": ["client_id_anonymized"],
}
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
BASE = ["supplier", "sku", "unit", "source", "provenance"]

def clean_header(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()

def number(value):
    if value is None or pd.isna(value) or str(value).strip() == "":
        return None
    try:
        n = float(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
        return n if np.isfinite(n) else None
    except ValueError:
        return None

def parse_date(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return pd.Timestamp(value).date()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None

def normalize_columns(frame):
    lookup = {a: k for k, aliases in ALIASES.items() for a in aliases}
    result = frame.rename(columns={c: lookup.get(clean_header(c), c) for c in frame}).copy()
    for c in ("sku", "unit"):
        if c in result:
            result[c] = result[c].astype("string")
    for c in ("quantity", "free_stock", "inbound_qty", "moq", "pack_multiple"):
        if c in result:
            result[c] = result[c].map(number)
    return result

def validate_table(frame, required):
    return [f"missing column: {c}" for c in required if c not in frame]

@dataclass
class Bundle:
    sales: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=BASE + ["date", "quantity", "document", "warehouse", "event_id"]))
    monthly: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=BASE + ["month", "quantity"]))
    stock: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=BASE + ["free_stock", "as_of"]))
    inbound: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=BASE + ["quantity", "eta", "shipment"]))
    policies: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(columns=BASE + ["moq", "pack_multiple", "category", "article"]))
    historical_stock: pd.DataFrame = field(default_factory=pd.DataFrame)
    warnings: list[str] = field(default_factory=list)
    manifest: list[dict] = field(default_factory=list)
    blocked_keys: set[tuple] = field(default_factory=set)
    snapshot: date = SNAPSHOT

def _files(name, content):
    if name.lower().endswith(".zip"):
        with ZipFile(BytesIO(content)) as archive:
            entries = [i for i in archive.infolist() if not i.is_dir() and i.filename.lower().endswith((".xlsx", ".csv"))]
            if len(entries) > 100 or sum(i.file_size for i in entries) > 300_000_000:
                raise ValueError("Archive exceeds 100 workbooks / 300 MB")
            for i in entries:
                yield f"{name}/{i.filename}", archive.read(i)
    else:
        yield name, content

def _sheets(name, content):
    if name.lower().endswith(".csv"):
        yield "CSV", csv.reader(StringIO(content.decode("utf-8-sig")))
    elif name.lower().endswith(".xlsx"):
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 400_000_000:
                raise ValueError("Workbook expanded size exceeds 400 MB")
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in book:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        raise ValueError("Supported files: ZIP, XLSX, CSV")

def _month_columns(headers):
    result = {}
    for i, raw in enumerate(headers):
        s = clean_header(raw)
        year = re.search(r"20\d{2}", s)
        if year:
            for m, prefix in enumerate(MONTHS, 1):
                if s.startswith(prefix):
                    result[i] = date(int(year[0]), m, 1)
            if re.fullmatch(r"20\d{2}-\d{2}", s):
                result[i] = date.fromisoformat(s + "-01")
    return result

def load_sources(sources, supplier=None, snapshot=SNAPSHOT):
    """Read paths or (filename, bytes, supplier) tuples without changing originals."""
    out = Bundle(snapshot=snapshot)
    tables = {k: [] for k in ("sales", "monthly", "stock", "inbound", "policies", "historical_stock")}
    seen = set()
    for source in sources:
        if isinstance(source, tuple):
            name, content, vendor = source
        else:
            path = Path(source)
            name, content = path.name, path.read_bytes()
            vendor = supplier or ("IEK" if "iek" in name.lower() else "Systeme Electric" if "systeme" in name.lower() else None)
        if vendor not in ("IEK", "Systeme Electric"):
            raise ValueError("Select IEK or Systeme Electric explicitly")
        for filename, data in _files(name, content):
            digest = sha256(data).hexdigest()
            if (vendor, digest) in seen:
                out.warnings.append(f"Repeated file ignored: {filename}")
                continue
            seen.add((vendor, digest))
            for sheet_name, iterator in _sheets(filename, data):
                mapping, count = None, 0
                line_counts = {}
                for row_no, row in enumerate(iterator, 1):
                    if mapping is None:
                        headers = list(row)
                        mapping = {target: next((i for i, c in enumerate(headers) if clean_header(c) in aliases), None) for target, aliases in ALIASES.items()}
                        if mapping["sku"] is None:
                            mapping = None
                            if row_no >= 20:
                                break
                            continue
                        months = _month_columns(headers)
                        shipments = {i: parse_date(re.findall(r"\d{2}\.\d{2}\.20\d{2}", str(h))[-1]) for i, h in enumerate(headers) if "поступление" in clean_header(h) and re.findall(r"\d{2}\.\d{2}\.20\d{2}", str(h))}
                        continue
                    def get(k):
                        i = mapping.get(k)
                        return row[i] if i is not None and i < len(row) else None
                    raw_sku = get("sku")
                    if raw_sku is None or str(raw_sku).strip().lower() in ("", "итого", "total"):
                        continue
                    sku = str(raw_sku).strip()
                    if not isinstance(raw_sku, str):
                        out.warnings.append(f"Numeric SKU: original leading zeroes unknown at {filename}:{row_no}")
                    ref = f"{filename}!{sheet_name}!{row_no}"
                    base = dict(supplier=vendor, sku=sku, unit=str(get("unit")).strip() if get("unit") else None, source=ref, provenance="actual")
                    count += 1
                    if mapping["date"] is not None and mapping["quantity"] is not None:
                        event = dict(base, date=parse_date(get("date")), quantity=number(get("quantity")), document=str(get("document") or ""), warehouse=str(get("warehouse") or ""), name=str(get("name") or ""), client_id_anonymized=get("client_id_anonymized"))
                        identity = f"{vendor}|{sku}|{get('date')}|{event['document']}|{event['warehouse']}|{base['unit']}|{event['quantity']}"
                        if not event["document"]:
                            identity += f"|unidentified:{ref}"
                        line_counts[identity] = line_counts.get(identity,0)+1
                        # Preserve repeated identical lines within one document/file.
                        # Deduplicate only matching occurrences in overlapping imports.
                        event["event_id"] = sha256(f"{identity}|occurrence:{line_counts[identity]}".encode()).hexdigest()[:24]
                        tables["sales"].append(event)
                    if months:
                        kind = "historical_stock" if "остат" in filename.lower() or "stock_history" in sheet_name.lower() else "monthly"
                        for i, month in months.items():
                            tables[kind].append(dict(base, month=month, quantity=number(row[i]), source=f"{ref}:{i+1}", partial=(month.year, month.month) == (snapshot.year, snapshot.month)))
                    if mapping["free_stock"] is not None:
                        stock_date = parse_date(get("as_of"))
                        tables["stock"].append(dict(base, free_stock=number(get("free_stock")), as_of=stock_date, reserved=number(get("reserved")), date_provenance="actual" if stock_date else "missing"))
                        if stock_date is None:
                            out.warnings.append(f"Undated stock: dated manager confirmation required at {ref}")
                    if any(mapping[k] is not None for k in ("moq", "pack_multiple", "category")):
                        tables["policies"].append(dict(base, moq=number(get("moq")), pack_multiple=number(get("pack_multiple")), category=str(get("category") or ""), article=str(get("article") or ""), ambiguous_moq=vendor == "IEK" and "мин. разр. к отгр." in [clean_header(x) for x in headers]))
                    if mapping["inbound_qty"] is not None:
                        tables["inbound"].append(dict(base, quantity=number(get("inbound_qty")), eta=parse_date(get("eta")), shipment=str(get("document") or ref)))
                    for i, eta in shipments.items():
                        if row[i] is not None:
                            tables["inbound"].append(dict(base, quantity=number(row[i]), eta=eta, shipment=str(headers[i]), source=f"{ref}:{i+1}"))
                out.manifest.append(dict(file=filename, sheet=sheet_name, sha256=digest, rows=count, recognized=mapping is not None))
                if mapping is None:
                    out.warnings.append(f"Reference/unrecognized sheet: {filename}!{sheet_name}")
    for kind, records in tables.items():
        if records:
            setattr(out, kind, pd.DataFrame(records))
    _audit(out)
    return out

def _audit(bundle):
    unit_sources = [t[KEY + ["unit"]] for t in (bundle.sales, bundle.stock, bundle.monthly, bundle.inbound, bundle.policies, bundle.historical_stock) if not t.empty and "unit" in t]
    unit_map = {}
    if unit_sources:
        known = pd.concat(unit_sources).dropna(subset=["unit"]).drop_duplicates()
        for key, group in known.groupby(KEY):
            units = group.unit.unique()
            if len(units) == 1:
                unit_map[key] = units[0]
            else:
                bundle.blocked_keys.add(key)
                bundle.warnings.append(f"Conflicting units: {key}: {list(units)}")
    for kind in ("sales", "monthly", "stock", "inbound", "policies", "historical_stock"):
        df = getattr(bundle, kind)
        if df.empty:
            continue
        df["unit_inferred"] = df.unit.isna()
        df["unit"] = [u if pd.notna(u) else unit_map.get((s, k)) for s, k, u in zip(df.supplier, df.sku, df.unit)]
        if kind == "sales":
            duplicates = df.duplicated("event_id", keep="first")
            if duplicates.any():
                bundle.warnings.append(f"Identical event repeats removed: {int(duplicates.sum())}; originals retained")
                df = df.loc[~duplicates].copy()
            ambiguous = df.duplicated(KEY + ["date", "document", "warehouse", "unit"], keep=False) & df.document.ne("")
            if ambiguous.any():
                bundle.warnings.append(f"Possible multi-line/conflicting documents retained for review: {int(ambiguous.sum())}")
            df["kind"] = np.select([df.quantity.isna(), df.quantity.lt(0), df.quantity.eq(0)], ["missing", "return", "zero"], default="sale")
        elif kind in ("stock", "policies", "monthly"):
            dims = KEY + (["month"] if kind == "monthly" else [])
            vals = [c for c in df if c not in {"source", "provenance", "date_provenance"}]
            df = df.drop_duplicates(vals)
            conflicts = df.duplicated(dims, keep=False)
            bundle.blocked_keys.update(df.loc[conflicts, KEY].itertuples(index=False, name=None))
            if conflicts.any():
                bundle.warnings.append(f"Conflicting {kind} rows blocked, not summed: {int(conflicts.sum())}")
        elif kind == "inbound":
            df = df.drop_duplicates(KEY + ["shipment", "eta", "quantity", "unit"])
        setattr(bundle, kind, df.reset_index(drop=True))
    if not bundle.sales.empty:
        if bundle.sales.date.isna().any():
            bundle.warnings.append(f"Sales rows with invalid/missing dates excluded from time calculations: {int(bundle.sales.date.isna().sum())}")
        for label, count in bundle.sales.kind.value_counts().items():
            bundle.warnings.append(f"Sales row kind / {label}: {count}")
    if "IEK" in set(bundle.sales.supplier) and "IEK" not in set(bundle.stock.supplier):
        bundle.warnings.append("IEK current stock missing; monthly opening balances are not current stock")
    bundle.warnings.append("Snapshot month incomplete; monthly blanks remain missing; no client identities used")

def reconcile(bundle):
    if bundle.monthly.empty or bundle.sales.empty:
        return pd.DataFrame()
    sales = bundle.sales.loc[bundle.sales.date.notna()].copy()
    sales["month"] = pd.to_datetime(sales.date).dt.to_period("M").dt.to_timestamp().dt.date
    trans = sales.groupby(KEY + ["month"], as_index=False).quantity.sum(min_count=1).rename(columns={"quantity": "transaction_net"})
    monthly = bundle.monthly.drop_duplicates(KEY + ["month"], keep=False)
    result = monthly.merge(trans, on=KEY + ["month"], how="outer", validate="one_to_one")
    result["difference"] = result.transaction_net - result.quantity
    return result
