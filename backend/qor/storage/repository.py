"""SQLite versioned workflow. Export reads only a persisted approved revision."""
from contextlib import contextmanager
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4
import csv
import json
import math
import sqlite3
from openpyxl import Workbook, load_workbook

TRANSITIONS = {"draft":"review", "review":"approved", "approved":"exported"}

def dump(value):
    return json.dumps(value,ensure_ascii=False,default=str,allow_nan=False)

def safe_cell(value):
    if isinstance(value,(dict,list)):
        value = dump(value)
    if isinstance(value,str) and (value.lstrip()[:1] in ("=","+","-","@") or value[:1] in ("\t","\r","\n")):
        return "'"+value
    return value

class Repository:
    def __init__(self,path):
        self.path = str(path)
        Path(path).parent.mkdir(parents=True,exist_ok=True)
        with self.db() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS revisions(run_id TEXT, version INTEGER, state TEXT, payload TEXT,
                context TEXT, actor TEXT, reason TEXT, created_at TEXT, PRIMARY KEY(run_id,version));
            CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, run_id TEXT, version INTEGER,
                action TEXT, actor TEXT, reason TEXT, at TEXT);
            CREATE TABLE IF NOT EXISTS inputs(id INTEGER PRIMARY KEY, dataset TEXT, kind TEXT, payload TEXT, at TEXT);
            ''')

    @contextmanager
    def db(self):
        con = sqlite3.connect(self.path,timeout=30)
        con.row_factory=sqlite3.Row
        try:
            with con:
                yield con
        finally:
            con.close()

    def record_input(self,dataset,kind,payload):
        with self.db() as db:
            db.execute("INSERT INTO inputs(dataset,kind,payload,at) VALUES(?,?,?,?)",(dataset,kind,dump(payload),datetime.now(timezone.utc).isoformat()))

    def inputs(self,dataset,kind):
        with self.db() as db:
            return [json.loads(r[0]) for r in db.execute("SELECT payload FROM inputs WHERE dataset=? AND kind=? ORDER BY id",(dataset,kind))]

    def _insert(self,db,run_id,version,state,rows,context,actor,reason):
        if not actor.strip():
            raise ValueError("Responsible manager is required")
        now=datetime.now(timezone.utc).isoformat()
        db.execute("INSERT INTO revisions VALUES(?,?,?,?,?,?,?,?)",(run_id,version,state,dump(rows),dump(context),actor,reason,now))
        db.execute("INSERT INTO audit(run_id,version,action,actor,reason,at) VALUES(?,?,?,?,?,?)",(run_id,version,state,actor,reason,now))

    def create(self,rows,context,actor):
        if not rows:
            raise ValueError("Empty draft")
        if len({r["supplier"] for r in rows}) != 1:
            raise ValueError("A draft must contain exactly one supplier")
        if len({(r["supplier"],r["sku"]) for r in rows}) != len(rows):
            raise ValueError("Duplicate SKU in draft")
        run_id=str(uuid4())
        normalized=[]
        for r in rows:
            normalized.append(dict(r,calculation_run_id=run_id,manager_qty=r.get("recommended_qty"),edit_reason=""))
        with self.db() as db:
            self._insert(db,run_id,1,"draft",normalized,context,actor,"Calculation created")
        return run_id

    def latest(self,run_id):
        with self.db() as db:
            r=db.execute("SELECT * FROM revisions WHERE run_id=? ORDER BY version DESC LIMIT 1",(run_id,)).fetchone()
        if r is None:
            raise ValueError("Unknown calculation run")
        result=dict(r)
        result["rows"]=json.loads(result.pop("payload"))
        result["context"]=json.loads(result["context"])
        return result

    def list_runs(self):
        with self.db() as db:
            return [dict(r) for r in db.execute("SELECT run_id,MAX(version) AS version,MAX(created_at) AS at FROM revisions GROUP BY run_id ORDER BY at DESC")]

    def edit(self,run_id,expected_version,edits,actor):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            old=self.latest(run_id)
            if old["version"]!=expected_version:
                raise ValueError("Draft changed: reload current version")
            rows=old["rows"]
            keyed={(r["supplier"],r["sku"]):r for r in rows}
            for edit in edits:
                key=(edit["supplier"],edit["sku"])
                if key not in keyed:
                    raise ValueError("Edit references unknown SKU")
                qty=float(edit["manager_qty"])
                reason=edit.get("reason", "").strip()
                if not math.isfinite(qty) or qty<0 or not reason:
                    raise ValueError("Finite non-negative quantity and edit reason required")
                if keyed[key].get("recommended_qty") is None:
                    raise ValueError("Incomplete item needs corrected data and a new calculation")
                keyed[key].update(manager_qty=qty,edit_reason=reason)
            self._insert(db,run_id,expected_version+1,"draft",rows,old["context"],actor,"Manager edits; approval reset")

    def transition(self,run_id,expected_version,target,actor):
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            old=self.latest(run_id)
            if old["version"]!=expected_version or TRANSITIONS.get(old["state"])!=target or target=="exported":
                raise ValueError("Invalid/stale transition; exported state requires generated verified artifact")
            if target=="approved":
                for r in old["rows"]:
                    q=r.get("manager_qty")
                    if q is None or not math.isfinite(float(q)) or q<0:
                        raise ValueError("Incomplete item prevents approval")
                    if r.get("urgency","").startswith("needs_") or r.get("urgency")=="data_conflict":
                        raise ValueError("Resolve source issues before approval")
            self._insert(db,run_id,expected_version+1,target,old["rows"],old["context"],actor,"Manager state transition")

    def export(self,run_id,expected_version,fmt,actor):
        if fmt not in ("csv","xlsx"):
            raise ValueError("Supported export formats: csv/xlsx")
        with self.db() as db:
            db.execute("BEGIN IMMEDIATE")
            old=self.latest(run_id)
            if old["version"]!=expected_version or old["state"] not in ("approved","exported"):
                raise ValueError("Export requires latest persisted approval")
            approval=db.execute("SELECT version,actor FROM revisions WHERE run_id=? AND state='approved' ORDER BY version DESC LIMIT 1",(run_id,)).fetchone()
            if approval is None:
                raise ValueError("Persisted approval missing")
            columns=["supplier","sku","article","unit","category","free_stock","inbound","daily_demand","forecast_method","lead_days","review_days","buffer_days","raw_need","binding_date","binding_projection","binding_buffer","moq","pack_multiple","recommended_qty","manager_qty","edit_reason","shortage_date","urgency","snapshot_date","stock_date","source_stock","source_policy","source_sales","forecast_history","candidate_count","assumptions","provenance","calculation_run_id","order_version","approval_status","approved_by"]
            data=[]
            for r in old["rows"]:
                row=dict(r,order_version=approval["version"],approval_status="approved",approved_by=approval["actor"])
                data.append([safe_cell(row.get(c)) for c in columns])
            if fmt=="csv":
                stream=StringIO(newline="")
                writer=csv.writer(stream); writer.writerow(columns); writer.writerows(data)
                payload=stream.getvalue().encode("utf-8-sig")
                reopened=list(csv.reader(StringIO(payload.decode("utf-8-sig"))))
                expected=[["" if x is None else str(x) for x in row] for row in data]
                if reopened[1:]!=expected:
                    raise ValueError("CSV verification failed")
            else:
                book=Workbook(); sheet=book.active; sheet.title="Approved order"
                sheet.append(columns)
                for row in data:
                    sheet.append(row)
                sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
                for cell in sheet["B"]:
                    cell.number_format="@"
                stream=BytesIO(); book.save(stream); payload=stream.getvalue()
                reopened=list(load_workbook(BytesIO(payload),data_only=False).active.values)[1:]
                for actual,expected in zip(reopened,data,strict=True):
                    for a,e in zip(actual,expected,strict=True):
                        if e=="" and a is None:
                            continue
                        if isinstance(e,(int,float)) and isinstance(a,(int,float)) and math.isclose(a,e,rel_tol=1e-12,abs_tol=1e-9):
                            continue
                        if a!=e:
                            raise ValueError("XLSX verification failed")
            self._insert(db,run_id,expected_version+1,"exported",old["rows"],old["context"],actor,f"Verified {fmt} artifact sha256={sha256(payload).hexdigest()}")
        return payload

    def audit(self,run_id):
        with self.db() as db:
            return [dict(r) for r in db.execute("SELECT * FROM audit WHERE run_id=? ORDER BY id",(run_id,))]
